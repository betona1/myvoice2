# -*- coding: utf-8 -*-
"""HSK 4~6급 전량 검증 엑셀. 검증자는 '수정뜻'과 '판정'만 채우면 된다."""
import os, sqlite3
os.chdir('/app')
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT='outputs/chinese/HSK4-6_단어검증.xlsx'
con=sqlite3.connect('file:database/voices.db?mode=ro',uri=True)
HEAD=Font(bold=True,color='FFFFFF',size=11); HFILL=PatternFill('solid',fgColor='4F46E5')
IN_FILL=PatternFill('solid',fgColor='FFF7CC'); T=Side(style='thin',color='D0D0D8')
BD=Border(left=T,right=T,top=T,bottom=T)
LVF={4:PatternFill('solid',fgColor='EAF4FF'),5:PatternFill('solid',fgColor='F1FBEF'),6:PatternFill('solid',fgColor='FFF1F1')}
COLS=[('ID',8),('급',5),('한자',12),('병음',18),('현재 뜻',34),('다의어(senses)',40),
      ('영어 뜻 (CC-CEDICT 원문)',46),('✏ 수정뜻',30),('✏ 판정',12),('비고',22)]
wb=Workbook(); ws=wb.active; ws.title='HSK4-6'
for i,(h,w) in enumerate(COLS,1):
    c=ws.cell(1,i,h); c.font=HEAD; c.fill=HFILL; c.border=BD
    c.alignment=Alignment(horizontal='center',vertical='center')
    ws.column_dimensions[get_column_letter(i)].width=w
ws.freeze_panes='C2'
rows=con.execute("""SELECT id,hsk,chinese,pinyin,meaning_ko,senses_ko,meaning_en
                    FROM words WHERE hsk BETWEEN 4 AND 6
                    ORDER BY hsk, freq DESC, id""").fetchall()
r=2
for wid,lv,ch,py,ko,sen,en in rows:
    vals=[wid,f'{lv}급',ch,py,ko or '',(sen or '') if (sen or '')!=(ko or '') else '',(en or '')[:400],'','','']
    for i,v in enumerate(vals,1):
        c=ws.cell(r,i,v); c.border=BD
        c.alignment=Alignment(vertical='center',wrap_text=(i in (5,6,7)))
        if i==2: c.fill=LVF[lv]; c.alignment=Alignment(horizontal='center',vertical='center')
        if i==3: c.font=Font(size=16,name='Malgun Gothic'); c.alignment=Alignment(horizontal='center',vertical='center')
        if i in (8,9,10): c.fill=IN_FILL
    ws.row_dimensions[r].height=26; r+=1
dv=DataValidation(type='list',formula1='"OK,수정,삭제,보류"',allow_blank=True)
dv.prompt='OK=정확 / 수정=오른쪽 수정뜻 기입 / 삭제=단어 아님 / 보류=판단 어려움'
dv.promptTitle='판정'; ws.add_data_validation(dv); dv.add(f'I2:I{r-1}')
ws.auto_filter.ref=f'A1:J{r-1}'

g=wb.create_sheet('작성요령')
guide=[['HSK 4~6급 단어 검증 요령',''],['',''],
 ['총 단어 수',f'{len(rows)}개  (4급 {sum(1 for x in rows if x[1]==4)} / 5급 {sum(1 for x in rows if x[1]==5)} / 6급 {sum(1 for x in rows if x[1]==6)})'],
 ['',''],
 ['① 확인 순서','「한자 + 병음 + 영어 뜻(CC-CEDICT)」을 보고 「현재 뜻」이 맞는지 판단'],
 ['','영어 뜻이 사전 원문이므로 가장 신뢰할 기준입니다.'],
 ['② 맞으면','「판정」에 OK (수정뜻은 비워둠)'],
 ['③ 틀리면','「판정」에 수정 + 「수정뜻」에 올바른 한국어 뜻'],
 ['④ 단어가 아니면','「판정」에 삭제'],
 ['⑤ 애매하면','「판정」에 보류 + 「비고」에 사유'],
 ['',''],
 ['뜻 표기 규칙','· 동사·형용사는 기본형으로: 「좁은」❌ → 「좁다」⭕'],
 ['','· 뜻이 여럿이면 쉼표로: 「들다, 제기하다」'],
 ['','· 양사는 용법을 괄호로: 「그루 (나무를 세는 양사)」'],
 ['','· 외래어보다 우리말 우선: 「베이스」❌ → 「기초, 토대」⭕'],
 ['',''],
 ['이미 자동교정한 것','병음 대문자 72 / 단일글자 뜻 311 / 양사·동사 큐레이션 82'],
 ['','관형형→기본형 315 / 문장형·사전잔재·오역 160'],
 ['남은 문제','기계번역 특유의 의미 오역 (예: 输入=수입하다 → 입력하다)'],
 ['','표본상 15~25% 추정. 이 파일로 걸러 주시면 DB에 일괄 반영합니다.'],
]
for i,(a,b) in enumerate(guide,1):
    g.cell(i,1,a).font=Font(bold=(i==1 or b==''),size=13 if i==1 else 11)
    g.cell(i,2,b).alignment=Alignment(wrap_text=True,vertical='center')
g.column_dimensions['A'].width=22; g.column_dimensions['B'].width=76
wb.move_sheet('작성요령',offset=-1)
os.makedirs('outputs/chinese',exist_ok=True); wb.save(OUT)
print(f"{OUT}  ({len(rows)}행)")
