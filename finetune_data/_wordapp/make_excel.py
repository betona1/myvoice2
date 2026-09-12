# -*- coding: utf-8 -*-
"""문제 단어 검증용 엑셀 생성.
검증자가 '수정뜻'과 '판정'만 채우면 되도록 구성한다."""
import os, json, sqlite3
os.chdir('/app')
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT='outputs/chinese/중국어_단어검증.xlsx'
probs=json.load(open('outputs/chinese/word_problems.json',encoding='utf-8'))
con=sqlite3.connect('file:database/voices.db?mode=ro',uri=True)
total=con.execute("SELECT COUNT(*) FROM words").fetchone()[0]

HEAD=Font(bold=True, color='FFFFFF', size=11)
HFILL=PatternFill('solid', fgColor='4F46E5')
IN_FILL=PatternFill('solid', fgColor='FFF7CC')     # 입력란
THIN=Side(style='thin', color='D0D0D8')
BD=Border(left=THIN,right=THIN,top=THIN,bottom=THIN)

wb=Workbook()

# ── 1) 안내 ──
ws=wb.active; ws.title='안내'
rows=[
 ['중국어 단어 검증표',''],
 ['',''],
 ['생성일', __import__('time').strftime('%Y-%m-%d %H:%M')],
 ['전체 단어', total],
 ['문제 단어', len(probs)],
 ['',''],
 ['작성 방법',''],
 ['1', '"검증" 시트에서 [현재뜻]이 맞는지 확인합니다.'],
 ['2', '틀렸으면 [수정뜻]에 올바른 한국어 뜻을 적습니다.'],
 ['3', '단어가 아니면(발음연습 조각 등) [판정]에 "삭제"를 고릅니다.'],
 ['4', '맞으면 [판정]에 "정상"을 고릅니다. 비워두면 미검토로 봅니다.'],
 ['5', '[영어뜻]은 CC-CEDICT 원문입니다. 판단 근거로 쓰세요.'],
 ['',''],
 ['주의', 'seq / id 열은 수정하지 마세요. 반영할 때 기준이 됩니다.'],
]
for r in rows: ws.append(r)
ws['A1'].font=Font(bold=True,size=16)
ws.column_dimensions['A'].width=14; ws.column_dimensions['B'].width=70

# ── 2) 유형별 요약 ──
ws2=wb.create_sheet('유형별요약')
from collections import Counter
cnt=Counter()
for p in probs:
    for t in p['problems'].split(';'): cnt[t]+=1
ws2.append(['문제유형','건수','설명'])
DESC={'사전미등재(장문)':'4자 이상인데 사전에 없음 — 두 단어가 붙었을 수 있음',
 '영어잔존':'한국어 뜻에 영어가 그대로 남음','뜻너무짧음':'단어는 긴데 뜻이 2자 이하',
 '뜻중복반복':'같은 말이 반복됨','사전표기잔존':'[병음]·변형·참조 등 사전 표기가 남음',
 '뜻잘림(수식어만)':'"매우"처럼 수식어만 남고 본뜻이 잘림','뜻없음':'뜻이 비어 있음',
 '병음섞임':'뜻에 병음이 섞임','설명문(단어아님)':'성조 설명 등 단어가 아닌 항목'}
for t,n in cnt.most_common():
    d=DESC.get(t, '다른 단어와 뜻이 똑같아 문제가 성립하지 않음' if t.startswith('뜻겹침') else '')
    ws2.append([t,n,d])
for c in ws2[1]: c.font=HEAD; c.fill=HFILL
ws2.column_dimensions['A'].width=22; ws2.column_dimensions['B'].width=8
ws2.column_dimensions['C'].width=60
ws2.freeze_panes='A2'

# ── 3) 검증 ──
ws3=wb.create_sheet('검증')
cols=['seq','id','한자','병음','현재뜻','수정뜻','판정','영어뜻','전체뜻','출처','급수','빈도','문제유형']
ws3.append(cols)
for p in probs:
    ws3.append([p['seq'],p['id'],p['chinese'],p['pinyin'],p['meaning_ko'],'','',
                p['meaning_en'],p['senses_ko'],p['subject'],p['level'],p['freq'],p['problems']])
for c in ws3[1]: c.font=HEAD; c.fill=HFILL; c.alignment=Alignment(horizontal='center')
W={'A':6,'B':7,'C':14,'D':20,'E':30,'F':30,'G':9,'H':46,'I':40,'J':16,'K':6,'L':7,'M':26}
for k,v in W.items(): ws3.column_dimensions[k].width=v
ws3.freeze_panes='D2'
ws3.auto_filter.ref=f"A1:M{ws3.max_row}"
# 입력란 강조 + 판정 드롭다운
dv=DataValidation(type='list', formula1='"정상,수정,삭제"', allow_blank=True)
ws3.add_data_validation(dv)
for r in range(2, ws3.max_row+1):
    ws3.cell(r,3).font=Font(name='SimSun', size=16, bold=True)   # 한자 크게
    ws3.cell(r,6).fill=IN_FILL; ws3.cell(r,7).fill=IN_FILL
    dv.add(ws3.cell(r,7))
    for cc in range(1,14):
        ws3.cell(r,cc).border=BD
        ws3.cell(r,cc).alignment=Alignment(vertical='center', wrap_text=(cc in (5,6,8,9)))

# ── 4) 유형별 분리 시트 (많은 것 위주) ──
for t,_ in cnt.most_common(6):
    name=t.replace('(','_').replace(')','')[:28]
    w=wb.create_sheet(name)
    w.append(cols)
    for p in probs:
        if t in p['problems'].split(';'):
            w.append([p['seq'],p['id'],p['chinese'],p['pinyin'],p['meaning_ko'],'','',
                      p['meaning_en'],p['senses_ko'],p['subject'],p['level'],p['freq'],p['problems']])
    for c in w[1]: c.font=HEAD; c.fill=HFILL
    for k,v in W.items(): w.column_dimensions[k].width=v
    w.freeze_panes='D2'; w.auto_filter.ref=f"A1:M{w.max_row}"
    for r in range(2, w.max_row+1):
        w.cell(r,3).font=Font(name='SimSun', size=16, bold=True)
        w.cell(r,6).fill=IN_FILL; w.cell(r,7).fill=IN_FILL

wb.save(OUT)
print(f"  저장: {OUT} ({os.path.getsize(OUT)//1024}KB)")
print(f"  시트: {', '.join(wb.sheetnames)}")
print(f"  검증 대상 {len(probs)}행")
