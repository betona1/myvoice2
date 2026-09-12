# -*- coding: utf-8 -*-
"""교정 엑셀을 DB에 반영.
 수정뜻 채워짐 → meaning_ko 교체 + senses_ko 재구성(교정 뜻을 앞에)
 판정=삭제     → 단어 및 학습기록 제거
 판정=정상     → 변경 없음
반영 후 seq(학습 순서)를 다시 매긴다."""
import os, re, sqlite3, sys
os.chdir('/app')
from openpyxl import load_workbook
DB='database/voices.db'
XL='/app/finetune_data/_wordapp/in/중국어_단어검증_해석교정완료.xlsx'
CIR='①②③④⑤⑥⑦⑧'

wb=load_workbook(XL, data_only=True)
ws=wb['검증']
hdr=[c.value for c in ws[1]]
ci={h:i for i,h in enumerate(hdr) if h}

fixes={}; kills=set(); okc=0
for r in ws.iter_rows(min_row=2, values_only=True):
    wid=r[ci['id']]
    if wid is None: continue
    f=(r[ci['수정뜻']] or '').strip() if r[ci['수정뜻']] else ''
    v=(r[ci['판정']] or '').strip() if r[ci['판정']] else ''
    if v=='삭제': kills.add(int(wid))
    elif f: fixes[int(wid)]=f
    elif v=='정상': okc+=1
print(f"  엑셀 판독 — 수정 {len(fixes)} / 삭제 {len(kills)} / 정상 {okc}")

con=sqlite3.connect(DB,timeout=60); con.execute('PRAGMA busy_timeout=60000')
cur=con.cursor()
have={r[0] for r in cur.execute("SELECT id FROM words")}
fixes={k:v for k,v in fixes.items() if k in have}
kills={k for k in kills if k in have}
print(f"  DB 대조 후 — 수정 {len(fixes)} / 삭제 {len(kills)}")

# 미리보기
prev=cur.execute(f"SELECT id,chinese,meaning_ko,senses_ko FROM words WHERE id IN ({','.join('?'*min(len(fixes),8))})",
                 list(fixes)[:8]).fetchall() if fixes else []
for wid,ch,ko,sen in prev:
    print(f"   {ch:8s} '{str(ko)[:22]}' → '{fixes[wid][:30]}'")

if '--commit' not in sys.argv:
    print("  [DRY RUN]"); con.close(); sys.exit()

# 1) 뜻 교정 — 교정 뜻을 대표로, 기존 뜻은 뒤에 남긴다
upd=[]
for wid, newko in fixes.items():
    row=cur.execute("SELECT senses_ko FROM words WHERE id=?", (wid,)).fetchone()
    old=[p.strip() for p in re.split(r'[①②③④⑤⑥⑦⑧]\s*', row[0] or '') if p.strip()]
    parts=[p.strip() for p in re.split(r'[;；]', newko) if p.strip()]   # "둘; 두" → ['둘','두']
    for o in old:
        if o not in parts: parts.append(o)
    sen=' '.join(f'{CIR[i]} {p}' for i,p in enumerate(parts[:6])) if len(parts)>1 else parts[0]
    upd.append((parts[0] if len(parts)==1 else newko, sen, wid))
cur.executemany("UPDATE words SET meaning_ko=?, senses_ko=? WHERE id=?", upd)
con.commit()
print(f"  ✅ 뜻 교정 {len(upd)}건")

# 2) 삭제
if kills:
    ids=list(kills); q=','.join('?'*len(ids))
    cur.execute(f"DELETE FROM word_progress WHERE word_id IN ({q})", ids)
    cur.execute(f"DELETE FROM study_log WHERE word_id IN ({q})", ids)
    cur.execute(f"DELETE FROM words WHERE id IN ({q})", ids)
    con.commit()
    print(f"  ✅ 삭제 {len(ids)}건")

# 3) 학습 순서 재부여
rows=cur.execute("SELECT id FROM words ORDER BY difficulty, freq DESC, chinese").fetchall()
cur.executemany("UPDATE words SET seq=? WHERE id=?", [(i+1, r[0]) for i,r in enumerate(rows)])
con.commit()
print(f"  ✅ seq 재부여 — 남은 단어 {len(rows)}개")
con.close()
